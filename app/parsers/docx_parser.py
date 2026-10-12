import docx

def parse_docx(path):
    chunks = []
    try:
        doc = docx.Document(path)
        
        # Headers/Footers
        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    if p.text.strip():
                        chunks.append({"text": p.text.strip(), "source": "header"})
            if section.footer:
                for p in section.footer.paragraphs:
                    if p.text.strip():
                        chunks.append({"text": p.text.strip(), "source": "footer"})

        # Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                chunks.append({"text": p.text.strip()})
                
        # Tables
        for table in doc.tables:
            table_rows_text = []
            headers = []
            for i, row in enumerate(table.rows):
                row_cells = [cell.text.strip() for cell in row.cells]
                if i == 0:
                    headers = row_cells
                    table_rows_text.append(" | ".join(headers))
                else:
                    table_rows_text.append(" | ".join(row_cells))
                    row_pairs = [
                        f"{headers[j] if j < len(headers) and headers[j] else f'Col{j}'}: {val}"
                        for j, val in enumerate(row_cells) if val
                    ]
                    if row_pairs:
                        chunks.append({"text": ", ".join(row_pairs), "source": "table"})
            if table_rows_text:
                chunks.append({"text": "\n".join(table_rows_text), "source": "table_full"})
    except Exception as e:
        print(f"Error parsing DOCX {path}: {e}")
    return chunks
