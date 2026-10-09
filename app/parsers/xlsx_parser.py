import openpyxl

def parse_xlsx(path):
    chunks = []
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c is not None else f"Col{j}" for j, c in enumerate(row)]
                else:
                    row_pairs = []
                    for j, cell in enumerate(row):
                        if cell is not None:
                            col_name = headers[j] if j < len(headers) else f"Col{j}"
                            row_pairs.append(f"{col_name}: {cell}")
                    if row_pairs:
                        chunks.append({"text": f"Sheet {sheet_name} Row {i}: " + ", ".join(row_pairs)})
    except Exception as e:
        print(f"Error parsing XLSX {path}: {e}")
    return chunks
